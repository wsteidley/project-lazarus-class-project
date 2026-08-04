"""Read the onshore-wind total-installed-cost series out of the IRENA RPGC workbook.

Python's half of the irena-tic extractor. It exists because the source is an .xlsx that
nothing on the TypeScript side reads. Its job is deliberately small: locate one row, get
the numbers out, and write a flat year,cost_usd_per_kw CSV. Everything about the derived
schema -- units, basis, segment, source stamps, rounding, sort order -- belongs to
src/lib/extractors/irena-tic.ts.

This is the Wright-fittable wind series. Wind LCOE is NOT (it folds in capacity-factor and
financing gains that aren't manufacturing learning, and fitting it yields a false
43%/doubling); total installed cost is the hardware curve.

Exit-code contract: 0 only when a plausible series was written; non-zero on any structural
surprise. A silently mis-aligned series would be committed as data -- which is not
hypothetical here: a hand-rolled zip+regex reader of this exact sheet shifted every year by
one column and produced a confident -87%/doubling. Hence openpyxl, and hence the guards
below locate rows by LABEL rather than by index.

Usage:
    uv run --project extract extract/irena_tic.py \
        --xlsx data/sources/irena/IRENA_TEC_RPGC_in_2025_data_file_2026.xlsx --out /tmp/tic.csv
"""

import argparse
import sys

from openpyxl import load_workbook

# Bind to the sheet TITLE, not the figure number: this series is "Fig 2.1" in the 2023
# edition and "Fig 2.3" in the 2025 one. The number is not stable across editions; the
# title is.
TITLE_MARKER = "tics of onshore wind projects"
FALLBACK_SHEET = "Fig 2.3"

# The series row. Matched EXACTLY (after strip/lower) rather than by substring, because the
# section header two rows above reads "Weighted average total installed costs (2025 USD/kW)"
# -- a substring match would find that first and read a row with no data.
SERIES_LABEL = "weighted average"

# The 5th/95th percentile rows describe the project RANGE, not the global series. They are
# deliberately not extracted: fitting a percentile band would answer a different question.
IGNORED_LABELS = ("5th percentile", "95th percentile")

# Currency guard. The TypeScript side stamps basis=real_2025_usd unconditionally, so a
# workbook denominated in anything else must stop the run rather than mislabel the series.
# Mixing a 2023-dollar series into 2025-dollar ones is the exact failure the one-basis-per-
# series rule exists to prevent.
REQUIRED_CURRENCY = "2025 usd"

FIRST_PLAUSIBLE_YEAR = 1990
LAST_PLAUSIBLE_YEAR = 2100
MIN_POINTS = 10
# Onshore wind TIC has run roughly $2,400/kW (2010) down to ~$980/kW (2025). This band is
# wide on purpose: it catches a row/column mis-read (which lands orders of magnitude out),
# not a data revision.
MIN_PLAUSIBLE_COST = 200.0
MAX_PLAUSIBLE_COST = 10_000.0


def find_sheet(workbook):
    """Locate the TIC sheet by its title text, falling back to the 2025 figure number."""
    for name in workbook.sheetnames:
        sheet = workbook[name]
        for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True):
            for value in row:
                if isinstance(value, str) and TITLE_MARKER in value.lower():
                    return sheet, name
    if FALLBACK_SHEET in workbook.sheetnames:
        return workbook[FALLBACK_SHEET], FALLBACK_SHEET
    raise SystemExit(
        f"no sheet whose first rows mention '{TITLE_MARKER}', and no '{FALLBACK_SHEET}' "
        f"sheet -- the workbook layout or figure numbering changed. Sheets: "
        f"{workbook.sheetnames[:20]}"
    )


def extract(xlsx_path: str) -> list[tuple[int, float]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet, sheet_name = find_sheet(workbook)
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    text = " ".join(
        str(cell).lower() for row in grid for cell in row if isinstance(cell, str)
    )
    if REQUIRED_CURRENCY not in text:
        raise SystemExit(
            f"sheet '{sheet_name}' does not state '{REQUIRED_CURRENCY}' anywhere -- this is "
            f"probably a different report edition. The derived rows are stamped "
            f"basis=real_2025_usd, so loading another currency basis would mislabel them."
        )

    # Year header: the row carrying the most plausible 4-digit years. Found by content, so
    # an inserted row above it cannot shift the series.
    year_row_index, years = None, {}
    for index, row in enumerate(grid):
        found = {
            position: int(cell)
            for position, cell in enumerate(row)
            if isinstance(cell, (int, float))
            and float(cell).is_integer()
            and FIRST_PLAUSIBLE_YEAR <= int(cell) <= LAST_PLAUSIBLE_YEAR
        }
        if len(found) > len(years):
            year_row_index, years = index, found
    if len(years) < MIN_POINTS:
        raise SystemExit(
            f"sheet '{sheet_name}': found only {len(years)} year columns (need "
            f"{MIN_POINTS}) -- the header row was not located"
        )

    # Series row: matched by label, below the year header.
    series_row = None
    for row in grid[year_row_index + 1 :]:
        labels = [
            cell.strip().lower() for cell in row if isinstance(cell, str) and cell.strip()
        ]
        if any(label in IGNORED_LABELS for label in labels):
            continue
        if any(label == SERIES_LABEL for label in labels):
            series_row = row
            break
    if series_row is None:
        raise SystemExit(
            f"sheet '{sheet_name}': no row labelled '{SERIES_LABEL}' below the year header "
            f"-- IRENA renamed the series row"
        )

    points: list[tuple[int, float]] = []
    for position, year in sorted(years.items()):
        cell = series_row[position] if position < len(series_row) else None
        if isinstance(cell, (int, float)):
            points.append((year, float(cell)))

    if len(points) < MIN_POINTS:
        raise SystemExit(
            f"sheet '{sheet_name}': only {len(points)} usable cost points (need "
            f"{MIN_POINTS}) -- the series row and the year header did not line up"
        )
    outliers = [
        (year, cost)
        for year, cost in points
        if not MIN_PLAUSIBLE_COST <= cost <= MAX_PLAUSIBLE_COST
    ]
    if outliers:
        raise SystemExit(
            f"sheet '{sheet_name}': cost values outside "
            f"{MIN_PLAUSIBLE_COST:,.0f}-{MAX_PLAUSIBLE_COST:,.0f} USD/kW: {outliers[:5]} "
            f"-- the wrong row or column range was read"
        )
    return points


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, help="path to the IRENA RPGC .xlsx")
    parser.add_argument("--out", required=True, help="path to write year,cost_usd_per_kw CSV")
    args = parser.parse_args()

    points = extract(args.xlsx)
    with open(args.out, "w", encoding="utf-8", newline="") as handle:
        handle.write("year,cost_usd_per_kw\n")
        for year, cost in points:
            handle.write(f"{year},{cost}\n")
    print(
        f"irena_tic: {len(points)} years ({points[0][0]}-{points[-1][0]}) -> {args.out}",
        file=sys.stderr,
    )


if __name__ == "__main__":
    main()
